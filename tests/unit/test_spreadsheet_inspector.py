from __future__ import annotations

import hashlib
import io
import xml.etree.ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from gaia_max.attachment_validation import AttachmentValidationError
from gaia_max.solvers.spreadsheet import WorkbookInspectionError, WorkbookInspector

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def audited_workbook_bytes() -> bytes:
    workbook = Workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.append(["Item", "Category", "Sales", "Units"])
    sales.append(["Burger", "Food", "12.50", 2])
    sales.append(["Soda", "Drink", "3.00", 1])
    sales["C4"] = "=SUM(C2:C3)"
    sales["C4"].number_format = "$#,##0.00"
    sales["C2"].number_format = "$#,##0.00"
    sales["C3"].number_format = "$#,##0.00"
    sales.row_dimensions[3].hidden = True
    sales.column_dimensions["D"].hidden = True
    sales["E1"] = "Notes"
    sales.merge_cells("E1:F1")
    sales.add_table(Table(displayName="SalesTable", ref="A1:C3"))
    sales.auto_filter.ref = "A1:C3"

    lookup = workbook.create_sheet("Lookup")
    lookup.sheet_state = "hidden"
    lookup.append(["Category", "Kind"])
    lookup.append(["Soda", "Drink"])
    workbook.defined_names.add(
        DefinedName("FoodItems", attr_text="'Sales'!$A$2:$A$2")
    )
    workbook.calculation.calcMode = "manual"
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = True

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_package(
    content: bytes,
    *,
    cached_formula_value: str | None = None,
    extra_parts: dict[str, bytes] | None = None,
) -> bytes:
    with ZipFile(io.BytesIO(content)) as source:
        parts = {name: source.read(name) for name in source.namelist()}

    if cached_formula_value is not None:
        root = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        for cell in root.findall(f".//{{{_SHEET_NS}}}c"):
            if cell.attrib.get("r") == "C4":
                value = cell.find(f"{{{_SHEET_NS}}}v")
                assert value is not None
                value.text = cached_formula_value
                break
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(
            root,
            encoding="utf-8",
            xml_declaration=True,
        )
    parts.update(extra_parts or {})

    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for name, payload in parts.items():
            target.writestr(name, payload)
    return output.getvalue()


def test_audit_identifies_all_relevant_workbook_structures() -> None:
    content = audited_workbook_bytes()

    audit = WorkbookInspector().inspect(content, original_name="sales.xlsx")

    assert audit.sha256 == hashlib.sha256(content).hexdigest()
    assert audit.sheet_names == ("Sales", "Lookup")
    assert audit.active_sheet_title == "Sales"
    assert audit.date_epoch == "1899-12-30"
    assert audit.defined_name_count == 1
    assert audit.hidden_sheet_count == 1
    assert audit.calculation_mode == "manual"
    assert audit.full_calculation_on_load is False
    assert audit.force_full_calculation is True

    sales = audit.sheets[0]
    assert sales.state == "visible"
    assert sales.dimensions == "A1:F4"
    assert sales.max_row == 4
    assert sales.max_column == 6
    assert sales.header_row == 1
    assert sales.headers == ("Item", "Category", "Sales", "Units", "Notes", "")
    assert sales.nonempty_cell_count == 14
    assert sales.hidden_rows == (3,)
    assert sales.hidden_columns == ("D",)
    assert sales.merged_ranges == ("E1:F1",)
    assert sales.tables == (("SalesTable", "A1:C3"),)
    assert sales.auto_filter_ref == "A1:C3"
    assert ("$#,##0.00", 3) in sales.number_formats
    assert sales.formula_count == 1
    assert sales.formula_cells[0].coordinate == "C4"
    assert sales.formula_cells[0].formula == "=SUM(C2:C3)"
    assert sales.formula_cells[0].cached_value_available is False
    assert audit.formula_count == 1
    assert audit.missing_formula_cache_count == 1


def test_formula_cache_presence_is_detected_without_exposing_cached_value() -> None:
    content = rewrite_package(audited_workbook_bytes(), cached_formula_value="15.5")

    audit = WorkbookInspector().inspect(content)

    formula = audit.sheets[0].formula_cells[0]
    assert formula.cached_value_available is True
    assert audit.missing_formula_cache_count == 0
    assert not hasattr(formula, "cached_value")


def test_package_risks_are_reported_even_when_parts_are_unreferenced() -> None:
    content = rewrite_package(
        audited_workbook_bytes(),
        extra_parts={
            "xl/externalLinks/externalLink1.xml": b"<externalLink/>",
            "xl/vbaProject.bin": b"synthetic macro bytes",
            "xl/calcChain.xml": b"<calcChain/>",
        },
    )

    audit = WorkbookInspector().inspect(content)

    assert audit.external_link_part_count == 1
    assert audit.has_vba_project is True
    assert audit.has_calculation_chain is True


def test_declared_grid_limit_blocks_sparse_dimension_bombs() -> None:
    workbook = Workbook()
    workbook.active.cell(row=2_000, column=2_000, value="far away")
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(WorkbookInspectionError, match="cell grid"):
        WorkbookInspector(max_audited_cells=1_000_000).inspect(output.getvalue())


def test_invalid_xlsx_is_rejected_before_openpyxl() -> None:
    with pytest.raises(AttachmentValidationError):
        WorkbookInspector().inspect(b"not an xlsx")


def test_structurally_minimal_package_that_openpyxl_cannot_load_fails_cleanly() -> None:
    output = io.BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        archive.writestr(
            "xl/workbook.xml",
            b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>',
        )

    with pytest.raises(WorkbookInspectionError, match="could not load"):
        WorkbookInspector().inspect(output.getvalue())
