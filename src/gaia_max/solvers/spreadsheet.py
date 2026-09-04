"""Answer-free structural inspection for validated XLSX workbooks."""

from __future__ import annotations

import asyncio
import hashlib
import io
import math
import re
from collections import Counter
from collections.abc import Awaitable, Callable, Iterable, Mapping
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from gaia_max.attachment_validation import (
    DEFAULT_MAX_ATTACHMENT_BYTES,
    AttachmentValidator,
)
from gaia_max.domain import AnswerType, DecimalAnswer, Modality, Question, TaskClass
from gaia_max.solver_routing import SolverRegistry, SolverRoute
from gaia_max.task_analysis import ReconciledTaskAnalysis

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DEFAULT_MAX_AUDITED_CELLS = 2_000_000
DEFAULT_MAX_WORKSHEETS = 100


class WorkbookInspectionError(ValueError):
    """A workbook cannot be audited completely and safely."""


class SpreadsheetCalculationError(ValueError):
    """Workbook rows cannot be classified or independently reproduced safely."""


@dataclass(frozen=True, slots=True)
class FormulaCellAudit:
    """One formula's location and cache status, without exposing its result value."""

    coordinate: str
    formula: str
    cached_value_available: bool
    number_format: str


@dataclass(frozen=True, slots=True)
class WorksheetAudit:
    """Structural facts needed before any spreadsheet calculation is trusted."""

    title: str
    state: str
    dimensions: str
    max_row: int
    max_column: int
    header_row: int | None
    headers: tuple[str, ...]
    nonempty_cell_count: int
    formula_cells: tuple[FormulaCellAudit, ...]
    number_formats: tuple[tuple[str, int], ...]
    hidden_rows: tuple[int, ...]
    hidden_columns: tuple[str, ...]
    merged_ranges: tuple[str, ...]
    tables: tuple[tuple[str, str], ...]
    auto_filter_ref: str | None

    @property
    def formula_count(self) -> int:
        return len(self.formula_cells)

    @property
    def missing_formula_cache_count(self) -> int:
        return sum(not formula.cached_value_available for formula in self.formula_cells)


@dataclass(frozen=True, slots=True)
class WorkbookAudit:
    """Complete answer-free workbook structure and calculation-risk report."""

    sha256: str
    sheet_names: tuple[str, ...]
    active_sheet_title: str
    date_epoch: str
    defined_name_count: int
    package_part_count: int
    external_link_part_count: int
    has_vba_project: bool
    has_calculation_chain: bool
    calculation_mode: str | None
    full_calculation_on_load: bool | None
    force_full_calculation: bool | None
    sheets: tuple[WorksheetAudit, ...]

    @property
    def formula_count(self) -> int:
        return sum(sheet.formula_count for sheet in self.sheets)

    @property
    def missing_formula_cache_count(self) -> int:
        return sum(sheet.missing_formula_cache_count for sheet in self.sheets)

    @property
    def hidden_sheet_count(self) -> int:
        return sum(sheet.state != "visible" for sheet in self.sheets)


class HiddenRowPolicy(StrEnum):
    """Explicit handling required when a data row is hidden."""

    ERROR = "error"
    INCLUDE = "include"
    EXCLUDE = "exclude"


class RowDisposition(StrEnum):
    """Auditable reason one non-empty row was included or excluded."""

    FOOD = "food"
    DRINK = "drink"
    TOTAL = "total"
    HIDDEN = "hidden"


@dataclass(frozen=True, slots=True)
class SpreadsheetCalculationPlan:
    """Reviewed workbook semantics; never inferred from a target answer."""

    sheet_name: str
    header_row: int
    item_column: str
    category_column: str
    sales_column: str
    included_categories: tuple[str, ...]
    excluded_categories: tuple[str, ...]
    total_markers: tuple[str, ...] = ("total", "subtotal", "grand total")
    hidden_row_policy: HiddenRowPolicy = HiddenRowPolicy.ERROR

    def __post_init__(self) -> None:
        columns = (self.item_column, self.category_column, self.sales_column)
        if not self.sheet_name or self.header_row < 1 or any(not column for column in columns):
            raise ValueError("spreadsheet plan sheet, row, and column fields must be populated")
        if len({column.casefold() for column in columns}) != len(columns):
            raise ValueError("spreadsheet plan columns must be distinct")
        included = _normalized_set(self.included_categories, field_name="included categories")
        excluded = _normalized_set(self.excluded_categories, field_name="excluded categories")
        _normalized_set(self.total_markers, field_name="total markers")
        if not included or not excluded:
            raise ValueError("spreadsheet plan requires included and excluded categories")
        if included.intersection(excluded):
            raise ValueError("included and excluded spreadsheet categories overlap")
        if not isinstance(self.hidden_row_policy, HiddenRowPolicy):
            raise ValueError("spreadsheet plan has an invalid hidden-row policy")


@dataclass(frozen=True, slots=True)
class ClassifiedSpreadsheetRow:
    """One normalized row and the explicit reason it was included or excluded."""

    excel_row: int
    item: str
    category: str
    sales: Decimal | None
    hidden: bool
    disposition: RowDisposition


@dataclass(frozen=True, slots=True)
class SpreadsheetPassResult:
    """One parser's exact total and classification audit."""

    engine: str
    total: Decimal
    rows: tuple[ClassifiedSpreadsheetRow, ...]

    @property
    def included_row_count(self) -> int:
        return sum(row.disposition is RowDisposition.FOOD for row in self.rows)


@dataclass(frozen=True, slots=True)
class SpreadsheetCalculationResult:
    """A total accepted only after independent parsers produce exact parity."""

    total: Decimal
    openpyxl_pass: SpreadsheetPassResult
    pandas_calamine_pass: SpreadsheetPassResult
    audit: WorkbookAudit


WorkbookSourceProvider = Callable[
    [Question, ReconciledTaskAnalysis],
    Awaitable[bytes],
]
SpreadsheetPlanProvider = Callable[
    [Question, ReconciledTaskAnalysis, WorkbookAudit],
    Awaitable[SpreadsheetCalculationPlan],
]


_TOTAL_LIKE = re.compile(r"\b(?:grand\s+total|sub\s*total|total)\b", re.IGNORECASE)
_CURRENCY_TEXT = re.compile(
    r"(?P<open>\()?\s*[$\u20ac\u00a3]?\s*"
    r"(?P<number>[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?)"
    r"\s*(?P<close>\))?"
)


def _normalized_set(values: tuple[str, ...], *, field_name: str) -> set[str]:
    normalized = {value.strip().casefold() for value in values}
    if "" in normalized or len(normalized) != len(values):
        raise ValueError(f"{field_name} contain empty or duplicate values")
    return normalized


def decimal_from_spreadsheet(value: object, *, field_name: str) -> Decimal:
    """Convert a workbook currency cell through decimal text, never binary arithmetic."""

    if isinstance(value, bool) or value is None:
        raise SpreadsheetCalculationError(f"{field_name} is not a currency value")
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, int):
        result = Decimal(value)
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise SpreadsheetCalculationError(f"{field_name} is not finite")
        result = Decimal(str(value))
    elif isinstance(value, str):
        text = value.strip()
        match = _CURRENCY_TEXT.fullmatch(text)
        if match is None:
            raise SpreadsheetCalculationError(f"{field_name} is not valid currency text")
        if bool(match.group("open")) != bool(match.group("close")):
            raise SpreadsheetCalculationError(f"{field_name} has unbalanced parentheses")
        normalized = match.group("number").replace(",", "")
        if match.group("open") and not normalized.startswith("-"):
            normalized = f"-{normalized}"
        try:
            result = Decimal(normalized)
        except InvalidOperation as exc:  # pragma: no cover - regex guards syntax
            raise SpreadsheetCalculationError(f"{field_name} is not valid decimal text") from exc
    else:
        raise SpreadsheetCalculationError(f"{field_name} has unsupported numeric type")
    if not result.is_finite():
        raise SpreadsheetCalculationError(f"{field_name} is not finite")
    return result


class WorkbookInspector:
    """Validate and inspect XLSX bytes without performing answer calculations."""

    def __init__(
        self,
        *,
        max_attachment_bytes: int = DEFAULT_MAX_ATTACHMENT_BYTES,
        max_audited_cells: int = DEFAULT_MAX_AUDITED_CELLS,
        max_worksheets: int = DEFAULT_MAX_WORKSHEETS,
    ) -> None:
        if max_audited_cells < 1 or max_worksheets < 1:
            raise ValueError("workbook inspection limits must be positive")
        self._validator = AttachmentValidator(max_bytes=max_attachment_bytes)
        self._max_audited_cells = max_audited_cells
        self._max_worksheets = max_worksheets

    def inspect(
        self,
        content: bytes,
        *,
        original_name: str = "workbook.xlsx",
        media_type: str = XLSX_MEDIA_TYPE,
    ) -> WorkbookAudit:
        """Return structural facts from formula and cached-value workbook views."""

        validation = self._validator.validate(
            content,
            original_name=original_name,
            media_type=media_type,
        )
        part_count, external_link_count, has_vba, has_calc_chain = self._package_facts(
            content
        )
        try:
            formula_book = load_workbook(
                io.BytesIO(content),
                data_only=False,
                read_only=False,
                keep_links=False,
            )
            cached_book = load_workbook(
                io.BytesIO(content),
                data_only=True,
                read_only=False,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise WorkbookInspectionError("openpyxl could not load the validated workbook") from exc

        try:
            if len(formula_book.worksheets) > self._max_worksheets:
                raise WorkbookInspectionError("workbook exceeds worksheet audit limit")
            if formula_book.sheetnames != cached_book.sheetnames:
                raise WorkbookInspectionError(
                    "formula and cached workbook views disagree on sheets"
                )

            remaining_cells = self._max_audited_cells
            sheet_audits: list[WorksheetAudit] = []
            for worksheet in formula_book.worksheets:
                declared_cells = worksheet.max_row * worksheet.max_column
                if declared_cells > remaining_cells:
                    raise WorkbookInspectionError("workbook declared cell grid exceeds audit limit")
                remaining_cells -= declared_cells
                cached_sheet = cached_book[worksheet.title]
                sheet_audits.append(self._inspect_sheet(worksheet, cached_sheet))

            calculation = formula_book.calculation
            active_sheet = formula_book.active
            if active_sheet is None:
                raise WorkbookInspectionError("workbook has no active worksheet")
            return WorkbookAudit(
                sha256=validation.sha256,
                sheet_names=tuple(formula_book.sheetnames),
                active_sheet_title=active_sheet.title,
                date_epoch=formula_book.epoch.date().isoformat(),
                defined_name_count=len(formula_book.defined_names),
                package_part_count=part_count,
                external_link_part_count=external_link_count,
                has_vba_project=has_vba,
                has_calculation_chain=has_calc_chain,
                calculation_mode=calculation.calcMode,
                full_calculation_on_load=calculation.fullCalcOnLoad,
                force_full_calculation=calculation.forceFullCalc,
                sheets=tuple(sheet_audits),
            )
        finally:
            formula_book.close()
            cached_book.close()

    @staticmethod
    def _package_facts(content: bytes) -> tuple[int, int, bool, bool]:
        with ZipFile(io.BytesIO(content)) as archive:
            names = tuple(archive.namelist())
        return (
            len(names),
            sum(
                name.startswith("xl/externalLinks/") and name.endswith(".xml")
                for name in names
            ),
            any(name.casefold().endswith("vbaproject.bin") for name in names),
            "xl/calcChain.xml" in names,
        )

    @classmethod
    def _inspect_sheet(
        cls,
        worksheet: Worksheet,
        cached_sheet: Worksheet,
    ) -> WorksheetAudit:
        nonempty_cells: list[Cell] = []
        formulas: list[FormulaCellAudit] = []
        formats: Counter[str] = Counter()
        first_nonempty_row: int | None = None

        for row in worksheet.iter_rows():
            row_has_value = False
            for cell in row:
                if cell.value is None:
                    continue
                row_has_value = True
                nonempty_cells.append(cell)
                formats[cell.number_format] += 1
                if cell.data_type == "f":
                    cached_value = cached_sheet[cell.coordinate].value
                    formulas.append(
                        FormulaCellAudit(
                            coordinate=cell.coordinate,
                            formula=str(cell.value),
                            cached_value_available=cached_value is not None,
                            number_format=cell.number_format,
                        )
                    )
            if row_has_value and first_nonempty_row is None:
                first_nonempty_row = row[0].row

        headers: tuple[str, ...] = ()
        if first_nonempty_row is not None:
            headers = tuple(
                "" if cell.value is None else str(cell.value)
                for cell in worksheet[first_nonempty_row]
            )

        table_audits = tuple(
            sorted((table.name, str(table.ref)) for table in worksheet.tables.values())
        )
        return WorksheetAudit(
            title=worksheet.title,
            state=worksheet.sheet_state,
            dimensions=worksheet.calculate_dimension(),
            max_row=worksheet.max_row,
            max_column=worksheet.max_column,
            header_row=first_nonempty_row,
            headers=headers,
            nonempty_cell_count=len(nonempty_cells),
            formula_cells=tuple(formulas),
            number_formats=tuple(sorted(formats.items())),
            hidden_rows=tuple(
                sorted(
                    index
                    for index, dimension in worksheet.row_dimensions.items()
                    if dimension.hidden
                )
            ),
            hidden_columns=tuple(
                sorted(
                    key
                    for key, dimension in worksheet.column_dimensions.items()
                    if dimension.hidden
                )
            ),
            merged_ranges=tuple(sorted(str(cell_range) for cell_range in worksheet.merged_cells)),
            tables=table_audits,
            auto_filter_ref=worksheet.auto_filter.ref,
        )


def _is_blank(value: object) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        return True
    try:
        missing = pd.isna(value)
    except (TypeError, ValueError):
        return False
    try:
        return bool(missing)
    except (TypeError, ValueError):
        return False


def _text_cell(value: object, *, field_name: str, required: bool = True) -> str:
    if _is_blank(value):
        if required:
            raise SpreadsheetCalculationError(f"{field_name} is blank")
        return ""
    if not isinstance(value, str):
        raise SpreadsheetCalculationError(f"{field_name} must contain text")
    text = value.strip()
    if required and not text:
        raise SpreadsheetCalculationError(f"{field_name} is blank")
    return text


def _column_indices(
    headers: Iterable[object],
    plan: SpreadsheetCalculationPlan,
) -> tuple[int, int, int]:
    positions: dict[str, int] = {}
    for index, raw_header in enumerate(headers):
        if _is_blank(raw_header):
            continue
        header = _text_cell(raw_header, field_name="header")
        normalized = header.casefold()
        if normalized in positions:
            raise SpreadsheetCalculationError(f"duplicate normalized header {header!r}")
        positions[normalized] = index

    requested = (plan.item_column, plan.category_column, plan.sales_column)
    try:
        found = tuple(positions[column.casefold()] for column in requested)
    except KeyError as exc:
        raise SpreadsheetCalculationError(f"required column is missing: {exc.args[0]}") from exc
    return found[0], found[1], found[2]


def _classify_records(
    records: Iterable[tuple[int, object, object, object]],
    *,
    hidden_rows: set[int],
    plan: SpreadsheetCalculationPlan,
    engine: str,
) -> SpreadsheetPassResult:
    included = _normalized_set(plan.included_categories, field_name="included categories")
    excluded = _normalized_set(plan.excluded_categories, field_name="excluded categories")
    total_markers = _normalized_set(plan.total_markers, field_name="total markers")
    classified: list[ClassifiedSpreadsheetRow] = []
    total = Decimal(0)

    for excel_row, raw_item, raw_category, raw_sales in records:
        if all(_is_blank(value) for value in (raw_item, raw_category, raw_sales)):
            continue
        item = _text_cell(raw_item, field_name=f"row {excel_row} item")
        category = _text_cell(
            raw_category,
            field_name=f"row {excel_row} category",
            required=False,
        )
        item_key = item.casefold()
        category_key = category.casefold()
        hidden = excel_row in hidden_rows

        if item_key in total_markers or category_key in total_markers:
            classified.append(
                ClassifiedSpreadsheetRow(
                    excel_row=excel_row,
                    item=item,
                    category=category,
                    sales=None,
                    hidden=hidden,
                    disposition=RowDisposition.TOTAL,
                )
            )
            continue
        if _TOTAL_LIKE.search(item) or _TOTAL_LIKE.search(category):
            raise SpreadsheetCalculationError(
                f"row {excel_row} looks like an unreviewed total row"
            )
        if not category:
            raise SpreadsheetCalculationError(f"row {excel_row} category is blank")
        sales = decimal_from_spreadsheet(raw_sales, field_name=f"row {excel_row} sales")
        if category_key in included:
            disposition = RowDisposition.FOOD
        elif category_key in excluded:
            disposition = RowDisposition.DRINK
        else:
            raise SpreadsheetCalculationError(
                f"row {excel_row} has unreviewed category {category!r}"
            )

        if hidden:
            if plan.hidden_row_policy is HiddenRowPolicy.ERROR:
                raise SpreadsheetCalculationError(
                    f"row {excel_row} is hidden and needs an explicit policy"
                )
            if plan.hidden_row_policy is HiddenRowPolicy.EXCLUDE:
                disposition = RowDisposition.HIDDEN
        row = ClassifiedSpreadsheetRow(
            excel_row=excel_row,
            item=item,
            category=category,
            sales=sales,
            hidden=hidden,
            disposition=disposition,
        )
        classified.append(row)
        if disposition is RowDisposition.FOOD:
            total += sales

    if not any(row.disposition is RowDisposition.FOOD for row in classified):
        raise SpreadsheetCalculationError("spreadsheet calculation selected no food rows")
    return SpreadsheetPassResult(engine=engine, total=total, rows=tuple(classified))


class SpreadsheetCalculator:
    """Recompute one reviewed workbook plan through independent parser engines."""

    def calculate(
        self,
        content: bytes,
        plan: SpreadsheetCalculationPlan,
        audit: WorkbookAudit,
    ) -> SpreadsheetCalculationResult:
        if hashlib.sha256(content).hexdigest() != audit.sha256:
            raise SpreadsheetCalculationError("workbook bytes do not match the supplied audit")
        if plan.sheet_name not in audit.sheet_names:
            raise SpreadsheetCalculationError("planned worksheet is missing from workbook audit")
        if audit.has_vba_project or audit.external_link_part_count:
            raise SpreadsheetCalculationError(
                "workbook has external or macro calculation dependencies"
            )

        openpyxl_result = self._calculate_openpyxl(content, plan)
        calamine_result = self._calculate_pandas_calamine(content, plan, audit)
        if openpyxl_result.total != calamine_result.total:
            raise SpreadsheetCalculationError("openpyxl and pandas totals disagree")
        if openpyxl_result.rows != calamine_result.rows:
            raise SpreadsheetCalculationError("openpyxl and pandas row classifications disagree")
        return SpreadsheetCalculationResult(
            total=openpyxl_result.total,
            openpyxl_pass=openpyxl_result,
            pandas_calamine_pass=calamine_result,
            audit=audit,
        )

    @staticmethod
    def _calculate_openpyxl(
        content: bytes,
        plan: SpreadsheetCalculationPlan,
    ) -> SpreadsheetPassResult:
        try:
            workbook = load_workbook(
                io.BytesIO(content),
                data_only=True,
                read_only=False,
                keep_links=False,
            )
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
            raise SpreadsheetCalculationError(
                "openpyxl calculation could not load workbook"
            ) from exc
        try:
            if plan.sheet_name not in workbook.sheetnames:
                raise SpreadsheetCalculationError("openpyxl calculation cannot find worksheet")
            sheet = workbook[plan.sheet_name]
            if plan.header_row > sheet.max_row:
                raise SpreadsheetCalculationError("planned header row is outside worksheet")
            indices = _column_indices(
                (cell.value for cell in sheet[plan.header_row]),
                plan,
            )
            hidden_rows = {
                index for index, dimension in sheet.row_dimensions.items() if dimension.hidden
            }
            records = (
                (
                    excel_row,
                    sheet.cell(excel_row, indices[0] + 1).value,
                    sheet.cell(excel_row, indices[1] + 1).value,
                    sheet.cell(excel_row, indices[2] + 1).value,
                )
                for excel_row in range(plan.header_row + 1, sheet.max_row + 1)
            )
            return _classify_records(
                records,
                hidden_rows=hidden_rows,
                plan=plan,
                engine="openpyxl",
            )
        finally:
            workbook.close()

    @staticmethod
    def _calculate_pandas_calamine(
        content: bytes,
        plan: SpreadsheetCalculationPlan,
        audit: WorkbookAudit,
    ) -> SpreadsheetPassResult:
        try:
            frames: Mapping[str, pd.DataFrame] = pd.read_excel(
                io.BytesIO(content),
                sheet_name=None,
                header=None,
                dtype=object,
                keep_default_na=False,
                na_filter=False,
                engine="calamine",
            )
        except (ImportError, OSError, TypeError, ValueError) as exc:
            raise SpreadsheetCalculationError("pandas Calamine could not load workbook") from exc
        if tuple(frames) != audit.sheet_names:
            raise SpreadsheetCalculationError(
                "pandas sheet inventory disagrees with workbook audit"
            )
        frame = frames[plan.sheet_name]
        if plan.header_row > len(frame.index):
            raise SpreadsheetCalculationError("planned header row is outside pandas worksheet")
        indices = _column_indices(frame.iloc[plan.header_row - 1].tolist(), plan)
        hidden_rows = set(
            next(sheet.hidden_rows for sheet in audit.sheets if sheet.title == plan.sheet_name)
        )
        records = (
            (
                position + 1,
                frame.iat[position, indices[0]],
                frame.iat[position, indices[1]],
                frame.iat[position, indices[2]],
            )
            for position in range(plan.header_row, len(frame.index))
        )
        return _classify_records(
            records,
            hidden_rows=hidden_rows,
            plan=plan,
            engine="pandas-calamine",
        )


class SpreadsheetSolver:
    """Resolve, audit, independently calculate, and return an exact Decimal answer."""

    def __init__(
        self,
        source_provider: WorkbookSourceProvider,
        plan_provider: SpreadsheetPlanProvider,
        *,
        inspector: WorkbookInspector | None = None,
        calculator: SpreadsheetCalculator | None = None,
    ) -> None:
        self._source_provider = source_provider
        self._plan_provider = plan_provider
        self._inspector = inspector or WorkbookInspector()
        self._calculator = calculator or SpreadsheetCalculator()

    async def solve(
        self,
        question: Question,
        analysis: ReconciledTaskAnalysis,
        route: str,
    ) -> object:
        if route != SolverRoute.SPREADSHEET.value:
            raise SpreadsheetCalculationError("spreadsheet solver received wrong route")
        if (
            analysis.task_class is not TaskClass.SPREADSHEET
            or analysis.modality is not Modality.XLSX
            or analysis.output_contract.answer_type is not AnswerType.CURRENCY
            or question.file_name is None
            or Path(question.file_name).suffix.casefold() != ".xlsx"
        ):
            raise SpreadsheetCalculationError(
                "spreadsheet solver received incompatible task metadata"
            )
        content = await self._source_provider(question, analysis)
        if not isinstance(content, bytes):
            raise SpreadsheetCalculationError("workbook source provider did not return bytes")
        audit = await asyncio.to_thread(
            self._inspector.inspect,
            content,
            original_name=question.file_name,
        )
        plan = await self._plan_provider(question, analysis, audit)
        if not isinstance(plan, SpreadsheetCalculationPlan):
            raise SpreadsheetCalculationError(
                "spreadsheet plan provider did not return a calculation plan"
            )
        result = await asyncio.to_thread(self._calculator.calculate, content, plan, audit)
        return DecimalAnswer(value=result.total)


def register_spreadsheet_solver(
    registry: SolverRegistry,
    source_provider: WorkbookSourceProvider,
    plan_provider: SpreadsheetPlanProvider,
    *,
    inspector: WorkbookInspector | None = None,
    calculator: SpreadsheetCalculator | None = None,
) -> SpreadsheetSolver:
    """Register B06 explicitly with answer-blind workbook and plan dependencies."""

    instance = SpreadsheetSolver(
        source_provider,
        plan_provider,
        inspector=inspector,
        calculator=calculator,
    )
    registry.register(SolverRoute.SPREADSHEET, instance)
    return instance
